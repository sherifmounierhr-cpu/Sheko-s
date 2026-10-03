import datetime as dt
from read_export import read_export
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule, CellIsRule
from openpyxl.comments import Comment

import sys

# Usage: python build_workbook.py <device-export.xls> <output.xlsx> [YYYY-MM-01]
SRC, OUT = sys.argv[1], sys.argv[2]
RAW = read_export(SRC)
if len(sys.argv) > 3:
    MONTH_START = dt.date.fromisoformat(sys.argv[3])
else:
    first = min(t for _, _, t in RAW)
    MONTH_START = dt.date(int(first[:4]), int(first[5:7]), 1)

EMP_SLOTS = 50
DAYS = 31
LOG_ROWS = 5000
PERM_ROWS = 300
MEET_ROWS = 300

F = "Arial"
NAVY = "1F3A5F"
HDR = PatternFill("solid", fgColor=NAVY)
INPUT = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

wb = Workbook()


def sheet(title, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def header(ws, row, cols, widths=None):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=11)
        c.fill = HDR
        c.alignment = CENTER
        c.border = BORDER
        if widths:
            ws.column_dimensions[L(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 34


def style(c, fill=None, fmt=None, bold=False, color="000000"):
    c.font = Font(name=F, bold=bold, color=color, size=10)
    c.alignment = CENTER
    c.border = BORDER
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def title(ws, text, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name=F, bold=True, size=15, color=NAVY)
    c.alignment = RIGHT
    ws.row_dimensions[1].height = 30


# ---------------------------------------------------------------- sheets
s_help = sheet("التعليمات", first=True)
s_set = sheet("الإعدادات")
s_emp = sheet("الموظفين")
s_log = sheet("سجل البصمة")
s_perm = sheet("الأذونات والإجازات")
s_meet = sheet("اجتماعات العملاء")
s_day = sheet("الحضور اليومي")
s_sum = sheet("الملخص الشهري")
s_lst = sheet("القوائم")

# ---------------------------------------------------------------- lists
DAYS_AR = ["الأحد", "الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت"]
lists = {
    "A": ("أيام الأسبوع", DAYS_AR + ["لا يوجد"]),
    "B": ("نوع الإذن", ["إذن تأخير", "إذن انصراف مبكر", "إذن خلال اليوم",
                        "إجازة اعتيادية", "إجازة عارضة", "إجازة مرضية", "إجازة بدون مرتب"]),
    "C": ("الحالة", ["معتمد", "قيد المراجعة", "مرفوض"]),
    "D": ("نوع المهمة", ["اجتماع مع عميل خارج المكتب", "معاينة موقع مع عميل",
                         "اجتماع مع عميل في المكتب", "مأمورية عمل", "معرض / إيفنت"]),
    "E": ("نتيجة الاجتماع", ["حجز / تعاقد", "مهتم", "متابعة لاحقة", "غير مهتم", "لم يحضر العميل"]),
    "F": ("مصدر العميل", ["فيسبوك", "إنستجرام", "جوجل", "تيك توك", "واتساب",
                          "ترشيح (ريفرال)", "مكالمة واردة", "زيارة مباشرة", "معرض"]),
}
for col, (name, vals) in lists.items():
    c = s_lst[f"{col}1"]
    c.value = name
    c.font = Font(name=F, bold=True)
    for i, v in enumerate(vals, 2):
        s_lst[f"{col}{i}"] = v
        s_lst[f"{col}{i}"].font = Font(name=F)
    s_lst.column_dimensions[col].width = 28
s_lst.sheet_state = "hidden"


def lst(col):
    n = len(lists[col][1])
    return f"القوائم!${col}$2:${col}${n + 1}"


def dv(ws, rng, src):
    d = DataValidation(type="list", formula1=f"={src}", allow_blank=True)
    d.error = "اختار من القائمة"
    d.errorTitle = "قيمة غير صحيحة"
    ws.add_data_validation(d)
    d.add(rng)


# ---------------------------------------------------------------- settings
ws = s_set
title(ws, "إعدادات نظام الحضور والانصراف - Everest Real Estate", 10)
for col, w in zip("ABCDEFGHIJ", [34, 16, 4, 18, 16, 4, 18, 16, 4, 20]):
    ws.column_dimensions[col].width = w

ws["A3"] = "الإعداد"
ws["B3"] = "القيمة"
for c in (ws["A3"], ws["B3"]):
    style(c, HDR, bold=True, color="FFFFFF")

settings = [
    ("أول يوم في الشهر المطلوب تحليله", MONTH_START, "dd/mm/yyyy",
     "غير التاريخ ده لأول يوم في الشهر اللي عايز تحلله، وكل الشيتات هتتحدث."),
    ("ميعاد الحضور الافتراضي", dt.time(11, 0), "hh:mm",
     "يتطبق على كل الموظفين، إلا اللي ليهم ميعاد خاص في شيت الموظفين."),
    ("ميعاد الانصراف الافتراضي", dt.time(19, 0), "hh:mm", "اتحدد 11 ص إلى 7 م حسب متوسط بصمات شهر سبتمبر. عدله لو المواعيد الرسمية مختلفة."),
    ("عدد أيام العمل لحساب قيمة اليوم", 26, "0",
     "قيمة اليوم = الراتب ÷ الرقم ده. شركات كتير بتستخدم 30 بدل 26."),
    ("الإجازة الأسبوعية (اليوم الأول)", "الجمعة", None, None),
    ("الإجازة الأسبوعية (اليوم الثاني)", "لا يوجد", None, "اختار (لا يوجد) لو الإجازة يوم واحد."),
    ("خصم البصمة الناقصة (يوم)", 0.5, "0.00", "لما الموظف يبصم مرة واحدة بس في اليوم."),
    ("خصم الغياب بدون إذن (يوم)", 1, "0.00", "بعض الشركات بتخصم يومين عن يوم الغياب."),
    ("رصيد الأذونات الشهري (ساعات)", 4, "0.0", "لو الموظف عدّاه، هيظهر تنبيه في الملخص."),
]
for i, (k, v, fmt, note) in enumerate(settings, 4):
    a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=v)
    style(a)
    a.alignment = RIGHT
    style(b, INPUT, fmt, color="0000FF")
    if note:
        b.comment = Comment(note, "Everest")
dv(ws, "B8:B9", lst("A"))
# helper weekday numbers (WEEKDAY type 1: Sunday = 1)
ws["C8"] = f'=IFERROR(MATCH(B8,{lst("A")},0),0)'
ws["C9"] = f'=IFERROR(MATCH(B9,{lst("A")},0),0)'
for r in (8, 9):
    ws[f"C{r}"].font = Font(name=F, color="FFFFFF", size=8)

SET = {
    "start": "الإعدادات!$B$4", "in": "الإعدادات!$B$5", "out": "الإعدادات!$B$6",
    "days": "الإعدادات!$B$7", "wk1": "الإعدادات!$C$8", "wk2": "الإعدادات!$C$9",
    "miss": "الإعدادات!$B$10", "abs": "الإعدادات!$B$11", "allow": "الإعدادات!$B$12",
}


def tier_table(col, label, rows, note):
    c1, c2 = col, chr(ord(col) + 1)
    ws.merge_cells(f"{c1}2:{c2}2")
    ws[f"{c1}2"] = label
    ws[f"{c1}2"].font = Font(name=F, bold=True, color=NAVY, size=11)
    ws[f"{c1}2"].alignment = CENTER
    ws[f"{c1}3"], ws[f"{c2}3"] = "من (دقيقة)", "الخصم (يوم)"
    style(ws[f"{c1}3"], HDR, bold=True, color="FFFFFF")
    style(ws[f"{c2}3"], HDR, bold=True, color="FFFFFF")
    for i, (m, d) in enumerate(rows, 4):
        style(ws.cell(row=i, column=ord(c1) - 64, value=m), INPUT, "0", color="0000FF")
        style(ws.cell(row=i, column=ord(c2) - 64, value=d), INPUT, "0.00", color="0000FF")
    ws[f"{c1}3"].comment = Comment(note, "Everest")
    return f"الإعدادات!${c1}$4:${c1}${3 + len(rows)}", f"الإعدادات!${c2}$4:${c2}${3 + len(rows)}"


LATE_M, LATE_D = tier_table("D", "شرائح خصم التأخير", [(0, 0), (16, 0.25), (31, 0.5), (61, 1)],
                            "أول 15 دقيقة سماح. خلي الأرقام تصاعدية، وتقدر تغيرها براحتك.")
EARLY_M, EARLY_D = tier_table("G", "شرائح خصم الانصراف المبكر", [(0, 0), (1, 0.25), (61, 0.5), (121, 1)],
                              "بدون سماح افتراضيًا. خلي الأرقام تصاعدية.")

ws["J2"] = "الإجازات الرسمية"
ws["J2"].font = Font(name=F, bold=True, color=NAVY, size=11)
ws["J2"].alignment = CENTER
ws["J3"] = "التاريخ"
style(ws["J3"], HDR, bold=True, color="FFFFFF")
holidays = [dt.date(2026, 10, 6)]
for r in range(4, 34):
    c = ws.cell(row=r, column=10, value=holidays[r - 4] if r - 4 < len(holidays) else None)
    style(c, INPUT, "dd/mm/yyyy", color="0000FF")
ws["J4"].comment = Comment("مثال: 6 أكتوبر. ضيف باقي الإجازات الرسمية تحتها.", "Everest")
HOL = "الإجازات!$J$4:$J$33".replace("الإجازات", "الإعدادات")

ws["A15"] = "الخانات الصفرا بس هي اللي بتتعدل. أي تغيير هنا بيتطبق فورًا على الحضور اليومي والملخص."
ws["A15"].font = Font(name=F, italic=True, color="7F7F7F")

# ---------------------------------------------------------------- employees
ws = s_emp
title(ws, "بيانات الموظفين", 9)
cols = ["كود البصمة", "الاسم", "القسم", "الوظيفة", "الراتب الأساسي (ج.م)",
        "ميعاد حضور خاص", "ميعاد انصراف خاص", "يوم إجازة خاص", "الموبايل", "الاسم في الجهاز", "ملاحظات"]
header(ws, 3, cols, [13, 26, 16, 20, 18, 15, 15, 15, 16, 20, 26])
import collections
seen = collections.OrderedDict()
for code, name, _ in RAW:
    seen.setdefault(code, name)
emps = sorted(((int(c), "" if n.strip() == c else n.strip().title(), n) for c, n in seen.items()), key=lambda x: x[0])
for r in range(4, 4 + EMP_SLOTS):
    e = emps[r - 4] if r - 4 < len(emps) else None
    vals = [e[0], e[1] or None, None, None, None, None, None, None, None, e[2]] if e else [None] * 10
    for ci, fmt in enumerate(["0", None, None, None, "#,##0", "hh:mm", "hh:mm", None, "@", None, None], 1):
        v = vals[ci - 1] if ci <= len(vals) else None
        cell = ws.cell(row=r, column=ci, value=v)
        if ci == 10:
            style(cell, CALC, None, color="7F7F7F")
        else:
            style(cell, INPUT, fmt, color="0000FF")
dv(ws, f"H4:H{3 + EMP_SLOTS}", lst("A"))
ws.conditional_formatting.add(f"B4:B{3 + EMP_SLOTS}", FormulaRule(
    formula=['AND($A4<>"",$B4="")'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws["B3"].comment = Comment("الخانات الحمرا = موظف متسجل في الجهاز من غير اسم. اكتب اسمه هنا وهيتحدث في كل الشيتات.", "Everest")
ws["H3"].comment = Comment("لو الموظف ليه يوم إجازة ثابت غير الإجازة العامة (زي يوم راحة في المبيعات) اختاره هنا.", "Everest")
ws["F3"].comment = Comment("سيبها فاضية لو الموظف على الميعاد الافتراضي في الإعدادات.", "Everest")
ws["A3"].comment = Comment("لازم يكون نفس الرقم المتسجل في جهاز البصمة بالظبط.", "Everest")
ws.freeze_panes = "C4"
E0, E1 = 4, 3 + EMP_SLOTS
EMP = lambda col: f"الموظفين!${col}${E0}:${col}${E1}"

# ---------------------------------------------------------------- raw log
ws = s_log
title(ws, "سجل البصمة الخام - الصق هنا التصدير من جهاز البصمة", 5)
header(ws, 3, ["AC-No. (الكود)", "Name (الاسم)", "Time (الوقت)", "تاريخ ووقت (تلقائي)", "اليوم (تلقائي)", "الكود (تلقائي)"],
       [14, 22, 22, 20, 14, 12])
ws["A3"].comment = Comment("الصق تصدير الجهاز زي ما هو (AC-No. / Name / Time) بداية من الخانة A4.", "Everest")

for r in range(4, 4 + LOG_ROWS):
    if r - 4 < len(RAW):
        code, name, t = RAW[r - 4]
        ws.cell(row=r, column=1, value=int(code))
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=t)
    ws.cell(row=r, column=4, value=(f'=IF(A{r}="","",IF(ISNUMBER(C{r}),C{r},IFERROR(DATE(LEFT(C{r},4),MID(C{r},6,2),MID(C{r},9,2))'
                                    f'+TIMEVALUE(TRIM(MID(C{r},12,20))),"")))')).number_format = "dd/mm/yyyy hh:mm"
    ws.cell(row=r, column=5, value=f'=IF(D{r}="","",INT(D{r}))').number_format = "dd/mm/yyyy"
    ws.cell(row=r, column=6, value=f'=IF(A{r}="","",TRIM(A{r}&""))')
ws.freeze_panes = "A4"
G0, G1 = 4, 3 + LOG_ROWS
LOG = lambda col: f"'سجل البصمة'!${col}${G0}:${col}${G1}"

# ---------------------------------------------------------------- permissions
ws = s_perm
title(ws, "الأذونات والإجازات", 13)
cols = ["كود البصمة", "الاسم (تلقائي)", "نوع الإذن / الإجازة", "من تاريخ", "إلى تاريخ (للإجازات)",
        "من الساعة", "إلى الساعة", "المدة بالدقائق (تلقائي)", "الحالة", "المعتمد من",
        "السبب / ملاحظات", "التصنيف (تلقائي)", "آخر يوم (تلقائي)"]
header(ws, 3, cols, [12, 24, 20, 13, 16, 11, 11, 15, 13, 16, 30, 12, 13])
perm_sample = [
    (55560, "إذن تأخير", dt.date(2026, 9, 10), None, dt.time(11, 0), dt.time(12, 30), "قيد المراجعة", "", "مثال - امسح الصف ده"),
]
for r in range(4, 4 + PERM_ROWS):
    v = perm_sample[r - 4] if r - 4 < len(perm_sample) else (None,) * 9
    vals = {1: v[0], 3: v[1], 4: v[2], 5: v[3], 6: v[4], 7: v[5], 9: v[6], 10: v[7], 11: v[8]}
    fmts = {1: "0", 4: "dd/mm/yyyy", 5: "dd/mm/yyyy", 6: "hh:mm", 7: "hh:mm"}
    for ci, val in vals.items():
        style(ws.cell(row=r, column=ci, value=val), INPUT, fmts.get(ci), color="0000FF")
    for ci, f, fmt in [
        (2, f'=IF(A{r}="","",IFERROR(INDEX({EMP("B")},MATCH(A{r},{EMP("A")},0)),"كود غير موجود"))', None),
        (8, f'=IF(OR(F{r}="",G{r}=""),0,MAX(0,ROUND((G{r}-F{r})*1440,0)))', "0"),
        (12, f'=IF(C{r}="","",IF(LEFT(C{r},5)="إجازة","إجازة","إذن"))', None),
        (13, f'=IF(D{r}="","",IF(E{r}="",D{r},E{r}))', "dd/mm/yyyy"),
    ]:
        style(ws.cell(row=r, column=ci, value=f), CALC, fmt)
dv(ws, f"C4:C{3 + PERM_ROWS}", lst("B"))
dv(ws, f"I4:I{3 + PERM_ROWS}", lst("C"))
ws.freeze_panes = "C4"
P1 = 3 + PERM_ROWS
PERM = lambda col: f"'الأذونات والإجازات'!${col}$4:${col}${P1}"

# ---------------------------------------------------------------- meetings
ws = s_meet
title(ws, "اجتماعات العملاء والمأموريات", 16)
cols = ["كود البصمة", "اسم الموظف (تلقائي)", "التاريخ", "من الساعة", "إلى الساعة", "المدة بالدقائق (تلقائي)",
        "نوع المهمة", "اسم العميل", "موبايل العميل", "المشروع", "مكان الاجتماع", "مصدر العميل",
        "نتيجة الاجتماع", "ميعاد المتابعة", "الحالة", "ملاحظات"]
header(ws, 3, cols, [12, 24, 13, 11, 11, 15, 26, 22, 16, 20, 22, 16, 16, 14, 13, 30])
meet_sample = [
    (55560, dt.date(2026, 9, 17), dt.time(11, 0), dt.time(12, 30), "اجتماع مع عميل خارج المكتب", "عميل (مثال)",
     "01200000000", "اسم المشروع", "كافيه - سموحة", "فيسبوك", "متابعة لاحقة",
     dt.date(2026, 9, 24), "قيد المراجعة", "مثال - امسح الصف ده"),
]
for r in range(4, 4 + MEET_ROWS):
    v = meet_sample[r - 4] if r - 4 < len(meet_sample) else (None,) * 14
    inputs = [1, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]
    fmts = {1: "0", 3: "dd/mm/yyyy", 4: "hh:mm", 5: "hh:mm", 9: "@", 14: "dd/mm/yyyy"}
    for ci, val in zip(inputs, v):
        style(ws.cell(row=r, column=ci, value=val), INPUT, fmts.get(ci), color="0000FF")
    style(ws.cell(row=r, column=2,
                  value=f'=IF(A{r}="","",IFERROR(INDEX({EMP("B")},MATCH(A{r},{EMP("A")},0)),"كود غير موجود"))'), CALC)
    style(ws.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),0,MAX(0,ROUND((E{r}-D{r})*1440,0)))'), CALC, "0")
dv(ws, f"G4:G{3 + MEET_ROWS}", lst("D"))
dv(ws, f"L4:L{3 + MEET_ROWS}", lst("F"))
dv(ws, f"M4:M{3 + MEET_ROWS}", lst("E"))
dv(ws, f"O4:O{3 + MEET_ROWS}", lst("C"))
ws.freeze_panes = "C4"
M1 = 3 + MEET_ROWS
MEET = lambda col: f"'اجتماعات العملاء'!${col}$4:${col}${M1}"

# ---------------------------------------------------------------- daily
ws = s_day
title(ws, "الحضور اليومي - يتحسب تلقائيًا (عدّل بس في عمود تعديل الخصم والملاحظات)", 26)
dcols = [
    ("k", "م", 4), ("d", "ي", 4), ("cnt", "ع", 4), ("code", "الكود", 9), ("name", "الاسم", 22), ("dept", "القسم", 13),
    ("date", "التاريخ", 12), ("wd", "اليوم", 10), ("rin", "الحضور المطلوب", 10), ("rout", "الانصراف المطلوب", 10),
    ("pin", "أول بصمة", 10), ("pout", "آخر بصمة", 10), ("hrs", "ساعات العمل", 9),
    ("type", "نوع اليوم", 12), ("leave", "إجازة معتمدة", 9), ("perm", "دقائق أذونات", 9),
    ("meet", "دقائق اجتماعات", 9), ("late", "دقائق تأخير", 9), ("early", "دقائق انصراف مبكر", 10),
    ("nlate", "صافي التأخير", 9), ("nearly", "صافي الانصراف المبكر", 10), ("status", "الحالة", 18),
    ("dlate", "خصم تأخير", 8), ("dearly", "خصم انصراف", 8), ("dmiss", "خصم بصمة ناقصة", 9),
    ("dabs", "خصم غياب", 8), ("dcalc", "الخصم المحسوب", 9), ("dman", "تعديل الخصم (يدوي)", 11),
    ("dfin", "الخصم النهائي (يوم)", 10), ("note", "ملاحظات", 26),
]
C = {k: L(i) for i, (k, _, _) in enumerate(dcols, 1)}
header(ws, 3, [n for _, n, _ in dcols], [w for _, _, w in dcols])
ws.row_dimensions[3].height = 46

D0 = 4
for idx in range(EMP_SLOTS * DAYS):
    r = D0 + idx
    k, d = idx // DAYS + 1, idx % DAYS
    c = {key: f"{C[key]}{r}" for key in C}
    er = E0 + k - 1
    blank = f'{c["date"]}=""'
    f = {
        "k": k, "d": d,
        "code": f'=IF(الموظفين!$A${er}="","",الموظفين!$A${er})',
        "name": f'=IF({c["code"]}="","",IF(الموظفين!$B${er}="","بدون اسم - "&{c["code"]},الموظفين!$B${er}))',
        "dept": f'=IF({c["code"]}="","",الموظفين!$C${er}&"")',
        "date": f'=IF({c["code"]}="","",IF(MONTH({SET["start"]}+{c["d"]})<>MONTH({SET["start"]}),"",{SET["start"]}+{c["d"]}))',
        "wd": f'=IF({blank},"",CHOOSE(WEEKDAY({c["date"]}),"الأحد","الاثنين","الثلاثاء","الأربعاء","الخميس","الجمعة","السبت"))',
        "rin": f'=IF({blank},"",IF(الموظفين!$F${er}="",{SET["in"]},الموظفين!$F${er}))',
        "rout": f'=IF({blank},"",IF(الموظفين!$G${er}="",{SET["out"]},الموظفين!$G${er}))',
        "cnt": f'=IF({blank},0,COUNTIFS({LOG("F")},{c["code"]}&"",{LOG("E")},{c["date"]}))',
        "pin": f'=IF({blank},"",IF({c["cnt"]}=0,"",_xlfn.MINIFS({LOG("D")},{LOG("F")},{c["code"]}&"",{LOG("E")},{c["date"]})-{c["date"]}))',
        "pout": f'=IF({blank},"",IF({c["cnt"]}<2,"",_xlfn.MAXIFS({LOG("D")},{LOG("F")},{c["code"]}&"",{LOG("E")},{c["date"]})-{c["date"]}))',
        "hrs": f'=IF(OR({c["pin"]}="",{c["pout"]}=""),"",({c["pout"]}-{c["pin"]})*24)',
        "type": f'=IF({blank},"",IF(OR(WEEKDAY({c["date"]})={SET["wk1"]},WEEKDAY({c["date"]})={SET["wk2"]},WEEKDAY({c["date"]})=IFERROR(MATCH(الموظفين!$H${er},{lst("A")},0),0)),"إجازة أسبوعية",IF(COUNTIF({HOL},{c["date"]})>0,"إجازة رسمية","يوم عمل")))',
        "leave": f'=IF({blank},"",COUNTIFS({PERM("A")},{c["code"]},{PERM("D")},"<="&{c["date"]},{PERM("M")},">="&{c["date"]},{PERM("I")},"معتمد",{PERM("L")},"إجازة"))',
        "perm": f'=IF({blank},"",SUMIFS({PERM("H")},{PERM("A")},{c["code"]},{PERM("D")},{c["date"]},{PERM("I")},"معتمد",{PERM("L")},"إذن"))',
        "meet": f'=IF({blank},"",SUMIFS({MEET("F")},{MEET("A")},{c["code"]},{MEET("C")},{c["date"]},{MEET("O")},"معتمد"))',
        "late": f'=IF(OR({blank},{c["pin"]}=""),"",MAX(0,ROUND(({c["pin"]}-{c["rin"]})*1440,0)))',
        "early": f'=IF(OR({blank},{c["pout"]}=""),"",MAX(0,ROUND(({c["rout"]}-{c["pout"]})*1440,0)))',
        "nlate": f'=IF({c["late"]}="","",MAX(0,{c["late"]}-{c["perm"]}-{c["meet"]}))',
        "nearly": f'=IF({c["early"]}="","",MAX(0,{c["early"]}-MAX(0,{c["perm"]}+{c["meet"]}-N({c["late"]}))))',
        "status": (f'=IF({blank},"",IF({c["type"]}<>"يوم عمل",{c["type"]},IF({c["leave"]}>0,"إجازة",'
                   f'IF({c["pin"]}="",IF({c["meet"]}>0,"مهمة خارجية","غياب"),'
                   f'IF({c["pout"]}="",IF({c["meet"]}>0,"حاضر - مهمة خارجية","بصمة ناقصة"),'
                   f'IF(AND({c["dlate"]}>0,{c["dearly"]}>0),"تأخير وانصراف مبكر",'
                   f'IF({c["dlate"]}>0,"تأخير",IF({c["dearly"]}>0,"انصراف مبكر","حاضر"))))))))'),
        "dlate": f'=IF({blank},"",IF(OR({c["type"]}<>"يوم عمل",{c["leave"]}>0,N({c["nlate"]})=0),0,LOOKUP({c["nlate"]},{LATE_M},{LATE_D})))',
        "dearly": f'=IF({blank},"",IF(OR({c["type"]}<>"يوم عمل",{c["leave"]}>0,N({c["nearly"]})=0),0,LOOKUP({c["nearly"]},{EARLY_M},{EARLY_D})))',
        "dmiss": f'=IF({blank},"",IF({c["status"]}="بصمة ناقصة",{SET["miss"]},0))',
        "dabs": f'=IF({blank},"",IF({c["status"]}="غياب",{SET["abs"]},0))',
        "dcalc": f'=IF({blank},"",{c["dlate"]}+{c["dearly"]}+{c["dmiss"]}+{c["dabs"]})',
        "dman": None,
        "dfin": f'=IF({blank},"",IF({c["dman"]}="",{c["dcalc"]},{c["dman"]}))',
        "note": None,
    }
    fmts = {"date": "dd/mm/yyyy", "rin": "hh:mm", "rout": "hh:mm", "pin": "hh:mm", "pout": "hh:mm",
            "hrs": "0.0", "dlate": "0.00;-0.00;-", "dearly": "0.00;-0.00;-", "dmiss": "0.00;-0.00;-",
            "dabs": "0.00;-0.00;-", "dcalc": "0.00;-0.00;-", "dman": "0.00", "dfin": "0.00;-0.00;-"}
    for key, val in f.items():
        cell = ws[c[key]]
        cell.value = val
        editable = key in ("dman", "note")
        style(cell, INPUT if editable else None, fmts.get(key), color="0000FF" if editable else "000000")
        cell.font = Font(name=F, size=9, color="0000FF" if editable else "000000",
                         bold=key in ("status", "dfin"))

DL = D0 + EMP_SLOTS * DAYS - 1
ws.column_dimensions["A"].hidden = True
ws.column_dimensions["B"].hidden = True
ws.column_dimensions["C"].hidden = True
ws.freeze_panes = f"{C['date']}{D0}"
ws.auto_filter.ref = f"A3:{C['note']}{DL}"
ws[C["dman"] + "3"].comment = Comment("اكتب هنا رقم الخصم اللي انت عايزه (0 = إلغاء الخصم). لو سبتها فاضية يتطبق الخصم المحسوب.", "Everest")

st = f"{C['status']}{D0}:{C['status']}{DL}"
for text, color in [("غياب", "F8CBAD"), ("بصمة ناقصة", "FFE699"), ("تأخير", "FCE4D6"),
                    ("انصراف مبكر", "FCE4D6"), ("إجازة", "DDEBF7"), ("مهمة خارجية", "E2EFDA"),
                    ("حاضر", "C6EFCE")]:
    ws.conditional_formatting.add(
        st, FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",{C["status"]}{D0}))'],
                        fill=PatternFill("solid", fgColor=color), stopIfTrue=True))
ws.conditional_formatting.add(
    f"{C['date']}{D0}:{C['note']}{DL}",
    FormulaRule(formula=[f'LEFT(${C["type"]}{D0},5)="إجازة"'], font=Font(color="808080")))

DAY = lambda key: f"'الحضور اليومي'!${C[key]}${D0}:${C[key]}${DL}"

# ---------------------------------------------------------------- summary
ws = s_sum
title(ws, "الملخص الشهري للحضور والخصومات", 24)
ws["A2"] = "الشهر:"
ws["B2"] = f'={SET["start"]}'
ws["B2"].number_format = "mmmm yyyy"
for a in ("A2", "B2"):
    ws[a].font = Font(name=F, bold=True, color=NAVY, size=12)

scols = [
    ("code", "الكود", 9), ("name", "الاسم", 24), ("dept", "القسم", 13), ("sal", "الراتب الأساسي", 12),
    ("req", "أيام العمل", 9), ("pres", "أيام الحضور", 9), ("abs", "غياب", 7), ("lv", "إجازات", 7),
    ("latec", "مرات التأخير", 9), ("latem", "دقائق التأخير", 9), ("earlyc", "مرات الانصراف المبكر", 10),
    ("miss", "بصمات ناقصة", 9), ("permh", "ساعات الأذونات", 9), ("permflag", "رصيد الأذونات", 11),
    ("meetc", "عدد الاجتماعات", 9), ("meeth", "ساعات الاجتماعات", 10), ("deals", "حجوزات من الاجتماعات", 10),
    ("commit", "نسبة الالتزام", 11), ("ddays", "أيام الخصم", 9), ("rate", "قيمة اليوم", 10),
    ("damt", "قيمة الخصم (ج.م)", 12), ("other", "خصومات أخرى", 11), ("bonus", "مكافآت / حوافز", 11),
    ("net", "صافي المرتب", 13),
]
S = {k: L(i) for i, (k, _, _) in enumerate(scols, 1)}
SH = 8
header(ws, SH, [n for _, n, _ in scols], [w for _, _, w in scols])
ws.row_dimensions[SH].height = 46
S0, S1 = SH + 1, SH + EMP_SLOTS
mstart, mend = SET["start"], f"EOMONTH({SET['start']},0)"
for i in range(EMP_SLOTS):
    r = S0 + i
    er = E0 + i
    c = {k: f"{S[k]}{r}" for k in S}
    code = c["code"]
    nb = f'{code}=""'
    f = {
        "code": f'=IF(الموظفين!$A${er}="","",الموظفين!$A${er})',
        "name": f'=IF({nb},"",IF(الموظفين!$B${er}="","بدون اسم - "&{code},الموظفين!$B${er}))',
        "dept": f'=IF({nb},"",الموظفين!$C${er}&"")',
        "sal": f'=IF({nb},"",N(الموظفين!$E${er}))',
        "req": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("type")},"يوم عمل"))',
        "abs": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("status")},"غياب"))',
        "lv": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("status")},"إجازة"))',
        "pres": f'=IF({nb},"",{c["req"]}-{c["abs"]}-{c["lv"]})',
        "latec": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("dlate")},">0"))',
        "latem": f'=IF({nb},"",SUMIFS({DAY("nlate")},{DAY("code")},{code}))',
        "earlyc": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("dearly")},">0"))',
        "miss": f'=IF({nb},"",COUNTIFS({DAY("code")},{code},{DAY("status")},"بصمة ناقصة"))',
        "permh": f'=IF({nb},"",SUMIFS({DAY("perm")},{DAY("code")},{code})/60)',
        "permflag": f'=IF({nb},"",IF({c["permh"]}>{SET["allow"]},"متجاوز","في الحدود"))',
        "meetc": f'=IF({nb},"",COUNTIFS({MEET("A")},{code},{MEET("O")},"معتمد",{MEET("C")},">="&{mstart},{MEET("C")},"<="&{mend}))',
        "meeth": f'=IF({nb},"",SUMIFS({MEET("F")},{MEET("A")},{code},{MEET("O")},"معتمد",{MEET("C")},">="&{mstart},{MEET("C")},"<="&{mend})/60)',
        "deals": f'=IF({nb},"",COUNTIFS({MEET("A")},{code},{MEET("O")},"معتمد",{MEET("M")},"حجز / تعاقد",{MEET("C")},">="&{mstart},{MEET("C")},"<="&{mend}))',
        "commit": f'=IF({nb},"",IF({c["req"]}-{c["lv"]}<=0,"",COUNTIFS({DAY("code")},{code},{DAY("status")},"حاضر*")+COUNTIFS({DAY("code")},{code},{DAY("status")},"مهمة خارجية"))/({c["req"]}-{c["lv"]})))',
        "ddays": f'=IF({nb},"",SUMIFS({DAY("dfin")},{DAY("code")},{code}))',
        "rate": f'=IF({nb},"",IF({SET["days"]}=0,0,{c["sal"]}/{SET["days"]}))',
        "damt": f'=IF({nb},"",ROUND({c["ddays"]}*{c["rate"]},2))',
        "other": None, "bonus": None,
        "net": f'=IF({nb},"",{c["sal"]}-{c["damt"]}-N({c["other"]})+N({c["bonus"]}))',
    }
    fmts = {"sal": "#,##0", "latem": "0", "permh": "0.0", "meeth": "0.0", "commit": "0%",
            "ddays": "0.00;-0.00;-", "rate": "#,##0.00", "damt": "#,##0;-#,##0;-", "other": "#,##0",
            "bonus": "#,##0", "net": "#,##0"}
    for key, val in f.items():
        editable = key in ("other", "bonus")
        style(ws[c[key]], INPUT if editable else None, fmts.get(key),
              bold=key in ("net", "commit"), color="0000FF" if editable else "000000")
        ws[c[key]].value = val

# totals row
tr = S1 + 1
ws.cell(row=tr, column=2, value="الإجمالي")
for key in ("sal", "abs", "lv", "latec", "latem", "earlyc", "miss", "permh", "meetc", "meeth", "deals",
            "ddays", "damt", "other", "bonus", "net"):
    ws[f"{S[key]}{tr}"] = f"=SUM({S[key]}{S0}:{S[key]}{S1})"
for ci in range(1, len(scols) + 1):
    cell = ws.cell(row=tr, column=ci)
    style(cell, PatternFill("solid", fgColor="D9E1F2"), ws.cell(row=S0, column=ci).number_format, bold=True)

ws.conditional_formatting.add(f"{S['commit']}{S0}:{S['commit']}{S1}",
                              DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                          color="5B9BD5"))
ws.conditional_formatting.add(f"{S['permflag']}{S0}:{S['permflag']}{S1}",
                              CellIsRule(operator="equal", formula=['"متجاوز"'],
                                         fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add(f"{S['abs']}{S0}:{S['abs']}{S1}",
                              CellIsRule(operator="greaterThanOrEqual", formula=["2"],
                                         fill=PatternFill("solid", fgColor="F8CBAD")))

# KPI cards
kpis = [
    ("عدد الموظفين", f'=COUNTIF({S["code"]}{S0}:{S["code"]}{S1},"<>")-COUNTIF({S["code"]}{S0}:{S["code"]}{S1},"")', "0"),
    ("متوسط نسبة الالتزام", f'=IFERROR(AVERAGE({S["commit"]}{S0}:{S["commit"]}{S1}),0)', "0%"),
    ("إجمالي أيام الغياب", f"={S['abs']}{tr}", "0"),
    ("إجمالي مرات التأخير", f"={S['latec']}{tr}", "0"),
    ("إجمالي الخصومات (ج.م)", f"={S['damt']}{tr}", "#,##0"),
    ("اجتماعات العملاء", f"={S['meetc']}{tr}", "0"),
    ("حجوزات من الاجتماعات", f"={S['deals']}{tr}", "0"),
]
col = 1
for label, formula, fmt in kpis:
    ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
    ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 2)
    a = ws.cell(row=4, column=col, value=label)
    a.font = Font(name=F, bold=True, color="FFFFFF", size=10)
    a.fill = HDR
    a.alignment = CENTER
    b = ws.cell(row=5, column=col, value=formula)
    b.font = Font(name=F, bold=True, color=NAVY, size=18)
    b.fill = PatternFill("solid", fgColor="EAF0F8")
    b.alignment = CENTER
    b.number_format = fmt
    col += 3
ws.freeze_panes = f"C{S0}"
ws["A{}".format(tr + 2)] = ("نسبة الالتزام = أيام الحضور في الميعاد ÷ (أيام العمل - الإجازات). "
                            "صافي المرتب = الراتب - قيمة الخصم - خصومات أخرى + مكافآت.")
ws["A{}".format(tr + 2)].font = Font(name=F, italic=True, color="7F7F7F")

# ---------------------------------------------------------------- help
ws = s_help
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 110
title(ws, "نظام الحضور والانصراف - Everest Real Estate", 2)
lines = [
    ("خطوات الاستخدام كل شهر", True),
    ("1. الإعدادات: حدد أول يوم في الشهر، ومواعيد الحضور والانصراف، والإجازة الأسبوعية، وشرائح الخصم والإجازات الرسمية.", False),
    ("2. الموظفين: الأكواد اتملت من الجهاز. الخانات الحمرا = موظفين من غير اسم، اكتب أسماءهم. ضيف القسم والراتب، وميعاد أو يوم إجازة خاص لو موجود.", False),
    ("3. سجل البصمة: فيه دلوقتي بيانات سبتمبر 2026. للشهر الجديد امسح الأعمدة A:C والصق تصدير الجهاز زي ما هو (AC-No. / Name / Time).", False),
    ("4. الأذونات والإجازات: سجل أي إذن أو إجازة، وخلي الحالة (معتمد) علشان تتحسب.", False),
    ("5. اجتماعات العملاء: سجل أي اجتماع أو معاينة أو مأمورية، وبيانات العميل ونتيجة الاجتماع. لو معتمد، الوقت ده مش هيتحسب تأخير أو غياب.", False),
    ("6. الحضور اليومي: بيتحسب لوحده. لو عايز تلغي أو تغير خصم يوم معين، اكتب الرقم في عمود (تعديل الخصم).", False),
    ("7. الملخص الشهري: الخصومات وصافي المرتب لكل موظف، وتقدر تضيف خصومات أخرى أو مكافآت.", False),
    ("", False),
    ("دليل الألوان", True),
    ("الخانات الصفرا بالخط الأزرق = بيانات انت بتدخلها وتعدلها.", False),
    ("الخانات الرمادي أو البيضا بالخط الأسود = معادلات تلقائية، متعدلهاش.", False),
    ("", False),
    ("قواعد الحساب", True),
    ("- التأخير = أول بصمة - ميعاد الحضور. الانصراف المبكر = ميعاد الانصراف - آخر بصمة.", False),
    ("- دقائق الأذونات والاجتماعات المعتمدة بتتخصم من التأخير الأول، والباقي منها من الانصراف المبكر.", False),
    ("- بصمة واحدة بس في اليوم = بصمة ناقصة، إلا لو عنده اجتماع أو مأمورية معتمدة في نفس اليوم.", False),
    ("- مفيش بصمة خالص + اجتماع أو مأمورية معتمدة = مهمة خارجية بدون خصم. من غيرها = غياب.", False),
    ("- الإجازة المعتمدة بتلغي كل خصومات اليوم.", False),
    ("", False),
    ("ملاحظات", True),
    ("- السعة: 50 موظف، و5,000 بصمة في الشهر، و300 إذن، و300 اجتماع. لو محتاج أكتر قولّي أكبّرها.", False),
    ("- الشيت بيقرا صيغة الوقت بتاعة الجهاز (2026-09-01 11:11 AM) تلقائيًا، وبيقرا كمان أي تاريخ ووقت حقيقي في Excel.", False),
    ("- صفوف (مثال) في الأذونات والاجتماعات حالتها (قيد المراجعة) فمش بتأثر على الحساب. امسحها قبل الاستخدام.", False),
]
for i, (t, bold) in enumerate(lines, 3):
    c = ws.cell(row=i, column=2, value=t)
    c.font = Font(name=F, bold=bold, size=12 if bold else 11, color=NAVY if bold else "000000")
    c.alignment = RIGHT

for w in wb.worksheets:
    w.sheet_properties.tabColor = {
        "التعليمات": "7F7F7F", "الإعدادات": "C00000", "الموظفين": "1F3A5F", "سجل البصمة": "1F3A5F",
        "الأذونات والإجازات": "BF8F00", "اجتماعات العملاء": "BF8F00", "الحضور اليومي": "548235",
        "الملخص الشهري": "548235"}.get(w.title, "FFFFFF")

wb.active = wb.sheetnames.index("الملخص الشهري")
wb.save(OUT)
print("saved", OUT)
